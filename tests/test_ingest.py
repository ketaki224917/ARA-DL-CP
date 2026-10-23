"""Tests for supported document loaders."""

from pathlib import Path

import pytest
from docx import Document as WordDocument

from ara.rag.ingest import IngestionError, load_directory, load_file


def test_load_text_and_markdown_files(tmp_path: Path) -> None:
    text_file = tmp_path / "plain.txt"
    markdown_file = tmp_path / "note.md"
    text_file.write_text("plain text", encoding="utf-8")
    markdown_file.write_text("# Heading\n\nBody", encoding="utf-8")

    assert load_file(text_file)[0].text == "plain text"
    assert load_file(markdown_file)[0].text == "# Heading\n\nBody"


def test_html_loader_removes_scripts_and_decodes_entities(tmp_path: Path) -> None:
    source = tmp_path / "page.html"
    source.write_text(
        "<html><body><h1>Research</h1><script>secret()</script><p>A &amp; B</p></body></html>",
        encoding="utf-8",
    )

    text = load_file(source)[0].text

    assert "Research" in text
    assert "A & B" in text
    assert "secret" not in text


def test_docx_loader_extracts_paragraphs_and_table_cells(tmp_path: Path) -> None:
    source = tmp_path / "note.docx"
    document = WordDocument()
    document.add_paragraph("Paragraph text")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.save(source)

    text = load_file(source)[0].text

    assert "Paragraph text" in text
    assert "left\tright" in text


def test_pdf_loader_preserves_page_numbers(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    class FakePage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class FakeReader:
        pages = [FakePage("first page"), FakePage("second page")]

        def __init__(self, _: str) -> None:
            pass

    source = tmp_path / "pages.pdf"
    source.write_bytes(b"mock")
    monkeypatch.setattr("ara.rag.ingest.PdfReader", FakeReader)

    documents = load_file(source)

    assert [(document.page, document.text) for document in documents] == [
        (1, "first page"),
        (2, "second page"),
    ]


def test_unsupported_file_and_invalid_directory_fail(tmp_path: Path) -> None:
    unsupported = tmp_path / "data.csv"
    unsupported.write_text("a,b", encoding="utf-8")

    with pytest.raises(IngestionError, match="Unsupported file type"):
        load_file(unsupported)
    with pytest.raises(IngestionError, match="not a directory"):
        load_directory(tmp_path / "missing")


def test_directory_load_is_sorted_and_uses_relative_source_names(tmp_path: Path) -> None:
    root = tmp_path / "corpus"
    root.mkdir()
    (root / "b.txt").write_text("second", encoding="utf-8")
    (root / "a.md").write_text("first", encoding="utf-8")
    (root / "skip.bin").write_bytes(b"binary")
    hidden = root / ".hidden"
    hidden.mkdir()
    (hidden / "secret.txt").write_text("ignored", encoding="utf-8")

    documents = load_directory(root)

    assert [(document.source, document.text) for document in documents] == [
        ("a.md", "first"),
        ("b.txt", "second"),
    ]

